#!/usr/bin/env python3
"""Presupuesto del Ramo 16 (Medio Ambiente y Recursos Naturales): aprobado,
modificado y pagado al segundo trimestre de 2026, y proyecto 2027.

Uso:
    python3 herramientas/extraer_presupuesto_ambiental.py AVANCE_2T_2026.csv PPEF_2027.xlsx

Fuentes (Transparencia Presupuestaria, Datos Abiertos):
  - /work/models/PTP/DatosAbiertos/Bases_de_datos_presupuesto/CSV/pef_ac01_avance_2t_2026.csv
  - /work/models/PTP/DatosAbiertos/Bases_de_datos_presupuesto/XLSX/PPEF_2027.xlsx
    (se usa el Excel: el CSV del proyecto 2027 que publica Hacienda viene
    cortado y sólo trae los tres primeros ramos).

Escribe investigaciones/presupuesto-ambiental.json; integrar_ambiente.py lo lee.
"""
import collections
import csv
import hashlib
import json
import pathlib
import sys

import openpyxl

RAIZ = pathlib.Path(__file__).resolve().parent.parent
SALIDA = RAIZ / 'investigaciones' / 'presupuesto-ambiental.json'
# Organos del sector que se muestran uno por uno; el resto de la SEMARNAT va junto.
UNIDADES = ['B00', 'RHQ', 'F00', 'E00', 'G00', 'RJJ']


def num(x):
    x = (x or '').strip().replace(',', '')
    return float(x) if x not in ('', '-') else 0.0


def sha(ruta):
    return hashlib.sha256(pathlib.Path(ruta).read_bytes()).hexdigest()


def avance(ruta):
    u = collections.OrderedDict()
    tot = [0.0, 0.0, 0.0]
    with open(ruta, encoding='utf-8-sig', newline='') as f:
        for r in csv.DictReader(f):
            if r['ID_RAMO'].strip() != '16':
                continue
            v = [num(r['MONTO_APROBADO']), num(r['MONTO_MODIFICADO']), num(r['MONTO_PAGADO'])]
            clave = r['ID_UR'].strip() if r['ID_UR'].strip() in UNIDADES else 'RESTO'
            d = u.setdefault(clave, {'ur': clave, 'nombre': r['DESC_UR'].strip() if clave != 'RESTO' else 'Resto de la SEMARNAT (oficinas centrales y representaciones)',
                                     'aprobado': 0.0, 'modificado': 0.0, 'pagado': 0.0})
            for i, k in enumerate(('aprobado', 'modificado', 'pagado')):
                d[k] += v[i]
                tot[i] += v[i]
    return u, dict(zip(('aprobado', 'modificado', 'pagado'), tot))


def proyecto(ruta):
    wb = openpyxl.load_workbook(ruta, read_only=True)
    ws = wb[wb.sheetnames[0]]
    filas = ws.iter_rows(values_only=True)
    cab = [str(c).strip() if c else '' for c in next(filas)]
    i = {c: cab.index(c) for c in cab}
    u = collections.defaultdict(float)
    nombres = {}
    total = 0.0
    for f in filas:
        if str(f[i['RAMO']]).strip() != '16':
            continue
        ur = str(f[i['UNIDAD']]).strip()
        clave = ur if ur in UNIDADES else 'RESTO'
        v = float(f[i['IMPORTE_PROYECTO']] or 0)
        u[clave] += v
        total += v
        nombres[ur] = f[i['UNIDAD_DESCRIPCION']]
    return u, total


def main():
    if len(sys.argv) != 3:
        sys.exit(__doc__)
    av, tot26 = avance(sys.argv[1])
    py, tot27 = proyecto(sys.argv[2])
    unidades = []
    for clave in UNIDADES + ['RESTO']:
        if clave not in av:
            continue
        d = {k: (round(v, 2) if isinstance(v, float) else v) for k, v in av[clave].items()}
        d['proyecto2027'] = round(py.get(clave, 0.0), 2)
        unidades.append(d)
    datos = {
        'ramo': '16', 'nombre': 'Medio Ambiente y Recursos Naturales',
        'total2026': {k: round(v, 2) for k, v in tot26.items()},
        'proyecto2027': round(tot27, 2),
        'unidades': unidades,
        'sha256': {'AV2T2026': sha(sys.argv[1]), 'PPEF2027': sha(sys.argv[2])},
    }
    SALIDA.write_text(json.dumps(datos, ensure_ascii=False, indent=1), encoding='utf-8')
    print(json.dumps(datos, ensure_ascii=False, indent=1))


if __name__ == '__main__':
    main()
